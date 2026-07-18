#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
figF_stress_strain_wtpct.py

Companion composition-series stress-strain figure:
Ecoflex 00-30 + NdFeB magnetic powder dogbone tension, 50/60/70/80 wt% NdFeB.

Inputs (Zwick exports):
  tensile/20241121/20241121/{50_1,60_1,70_2,80_5}.xlsx
  - sheet "测试结果": per-specimen Et / sm / em / A0 (two header rows, col0 "试样 N")
  - specimen sheets "试样 N": columns [应变 %, 标准载荷 <unit>]
    NOTE: in these exports the "load" column unit row reads "kPa", i.e. the
    machine already divided by A0 -> the column IS stress. The unit row is read
    at runtime; if a sheet were exported in N the script divides by that
    specimen's A0 instead. Nothing is hardcoded.

Outputs:
  paper/figures/fig2d_stress_strain_wtpct.{svg,pdf,png}

Style: materials-journal convention -- serif (Times New Roman / STIX), white
background, no grid, four-sided box, inward ticks on all sides, light raw
per-specimen curves + bold group mean, legend inside lower right with
Et = mean ± s.d. computed from the summary sheet at runtime.
"""

import os
import numpy as np
import openpyxl
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib import font_manager

DATA_DIR = ("/private/tmp/claude-502/-Users-arielzhang-Desktop-SensingMagic/"
            "617d7206-542d-4a74-bd36-6e9400bc4ca0/scratchpad/tensile/20241121/20241121")
FIG_DIR = "/Users/arielzhang/Desktop/SensingMagic/paper/figures"
BASENAME = "fig2d_stress_strain_wtpct"

# colorblind-safe ordered ramp (Paul Tol bright hues, low->high filler loading)
GROUPS = [  # (file, wt%, color)
    ("50_1.xlsx", 50, "#CCBB44"),
    ("60_1.xlsx", 60, "#66CCEE"),
    ("70_2.xlsx", 70, "#4477AA"),
    ("80_5.xlsx", 80, "#AA3377"),
]

SUMMARY_SHEET = "测试结果"
N_GRID = 900          # resample points per curve (visual only)
SMOOTH_WIN = 41       # mild moving-average window on the raw trace (visual only)


# ----------------------------------------------------------------------------- fonts
def setup_style():
    available = {f.name for f in font_manager.fontManager.ttflist}
    serif_stack = ["STIXGeneral", "STIX Two Text", "DejaVu Serif"]
    if "Times New Roman" in available:
        serif_stack.insert(0, "Times New Roman")
    plt.rcParams.update({
        "font.family": "serif",
        "font.serif": serif_stack,
        "mathtext.fontset": "stix",
        "figure.facecolor": "white",
        "axes.facecolor": "white",
        "axes.grid": False,
        "axes.linewidth": 0.9,
        "xtick.direction": "in",
        "ytick.direction": "in",
        "svg.fonttype": "none",
        "pdf.fonttype": 42,
    })


# ----------------------------------------------------------------------------- io
def read_summary(wb):
    """Parse the 测试结果 sheet -> {specimen_name: {Et, sm, em, A0, ...}}."""
    ws = wb[SUMMARY_SHEET]
    rows = [r for r in ws.iter_rows(values_only=True)]
    header = rows[0]  # (None, 'Et', 'sm', 'em', 'sb', 'eb', 'b', 'h', 'A0')
    col = {str(h): i for i, h in enumerate(header) if h is not None}
    out = {}
    for r in rows[2:]:
        if r is None or r[0] is None:
            continue
        name = str(r[0]).strip()
        if not name.startswith("试样"):
            continue
        out[name] = {k: float(r[i]) for k, i in col.items()
                     if r[i] is not None}
    return out


def read_specimen(ws, A0):
    """Read one 试样 sheet -> (strain %, stress MPa). Unit row decides scaling."""
    it = ws.iter_rows(values_only=True)
    next(it)                      # row 0: ('试样 N', '试样 N')
    next(it)                      # row 1: ('应变', '标准载荷')
    units = next(it)              # row 2: ('%', 'kPa' | 'N' | 'MPa')
    u = str(units[1]).strip().lower()
    strain, load = [], []
    for r in it:
        if r[0] is None or r[1] is None:
            continue
        strain.append(float(r[0]))
        load.append(float(r[1]))
    strain = np.asarray(strain)
    load = np.asarray(load)
    if u == "kpa":                # machine already normalized by A0
        stress = load / 1000.0
    elif u == "mpa":
        stress = load
    elif u == "n":                # true load -> engineering stress
        stress = load / A0
    else:
        raise ValueError(f"unknown load unit {u!r}")
    keep = strain >= 0.0
    return strain[keep], stress[keep]


def resample(strain, stress, n=N_GRID, win=SMOOTH_WIN):
    """Mild smoothing + downsample onto a uniform strain grid (visual only)."""
    if win > 1 and len(stress) > 3 * win:
        k = np.ones(win) / win
        stress = np.convolve(stress, k, mode="same")
        # fix edge bias of 'same' convolution
        stress[:win] = stress[win]
        stress[-win:] = stress[-win - 1]
    order = np.argsort(strain, kind="stable")
    s_sorted, y_sorted = strain[order], stress[order]
    grid = np.linspace(0.0, s_sorted[-1], n)
    return grid, np.interp(grid, s_sorted, y_sorted)


# ----------------------------------------------------------------------------- main
def main():
    setup_style()
    os.makedirs(FIG_DIR, exist_ok=True)

    fig, ax = plt.subplots(figsize=(4.6, 3.6), dpi=300)

    handles, labels = [], []
    print("group summary (computed at runtime):")
    for fname, wt, color in GROUPS:
        wb = openpyxl.load_workbook(os.path.join(DATA_DIR, fname),
                                    read_only=True, data_only=True)
        summary = read_summary(wb)
        spec_sheets = [s for s in wb.sheetnames if s.strip().startswith("试样")]

        curves = []
        for sheet in spec_sheets:
            A0 = summary.get(sheet.strip(), {}).get("A0", np.nan)
            s, y = read_specimen(wb[sheet], A0)
            curves.append(resample(s, y))
        wb.close()

        # light per-specimen raw curves
        for g, y in curves:
            ax.plot(g, y, color=color, lw=0.6, alpha=0.35, zorder=2)

        # bold group mean, only where ALL specimens are still intact
        cutoff = min(g[-1] for g, _ in curves)
        grid = np.linspace(0.0, cutoff, N_GRID)
        mean = np.mean([np.interp(grid, g, y) for g, y in curves], axis=0)
        ax.plot(grid, mean, color=color, lw=1.8, zorder=3)

        et = np.array([v["Et"] for v in summary.values()])
        et_m, et_s = et.mean(), et.std(ddof=1)
        print(f"  {wt} wt%: n={len(curves)}, Et = {et_m:.3f} ± {et_s:.3f} MPa, "
              f"mean cut at {cutoff:.0f}% strain")

        handles.append(Line2D([0], [0], color=color, lw=1.8))
        labels.append(f"{wt} wt%  $E_t$ = {et_m:.3f} $\\pm$ {et_s:.3f} MPa")

    ax.set_xlabel("Strain (%)", fontsize=11)
    ax.set_ylabel("Stress (MPa)", fontsize=11)
    ax.set_xlim(0, 800)
    ax.set_ylim(0, 1.25)
    ax.tick_params(direction="in", top=True, right=True, labelsize=9.5,
                   length=3.5, width=0.8)
    for sp in ax.spines.values():
        sp.set_visible(True)

    leg = ax.legend(handles, labels, loc="lower right", fontsize=8.3,
                    frameon=True, framealpha=1.0, edgecolor="0.7",
                    borderpad=0.6, handlelength=1.6, handletextpad=0.7,
                    labelspacing=0.45)
    leg.get_frame().set_linewidth(0.6)

    fig.tight_layout(pad=0.4)
    for ext in ("svg", "pdf", "png"):
        fig.savefig(os.path.join(FIG_DIR, f"{BASENAME}.{ext}"),
                    dpi=300, facecolor="white")
    plt.close(fig)
    print("saved:", os.path.join(FIG_DIR, BASENAME) + ".{svg,pdf,png}")


if __name__ == "__main__":
    main()
